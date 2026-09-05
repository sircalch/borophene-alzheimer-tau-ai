"""
prepare_tau_final_submission_package.py
Packages all generated files for Article 4 (Alzheimer's Tau & 2D Borophene) into an official
submission-ready folder and ZIP package.
"""

import os
import shutil
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def create_tau_cover_letter(sub_dir):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(33, 33, 33)
    
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_after = Pt(14)
    p_h.add_run(
        "Andrés Monreal Hernández, Ph.D.\n"
        "Universidad Estatal de Sonora\n"
        "Hermosillo, Sonora, Mexico\n"
        "Email: andres.monreal@ues.mx | ORCID: 0009-0009-1207-8597\n"
        "Date: August 30, 2026\n"
    ).font.bold = True
    
    p_ed = doc.add_paragraph()
    p_ed.paragraph_format.space_after = Pt(12)
    p_ed.add_run(
        "To: The Editor-in-Chief\n"
        "Beilstein Journal of Nanotechnology\n"
        "Beilstein-Institut, Frankfurt am Main, Germany\n"
    )
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    p_sub.add_run("Subject: Submission of Original Research Article for Peer Review").font.bold = True
    
    doc.add_paragraph("Dear Editor-in-Chief,")
    doc.add_paragraph(
        "On behalf of my co-authors (Sara Lizbeth Franco Amaya, Carlos Ivanhoe Martínez Osorio, and myself), "
        "I am pleased to submit our original research manuscript titled:"
    )
    
    p_t = doc.add_paragraph()
    p_t.paragraph_format.left_indent = Inches(0.4)
    p_t.paragraph_format.space_after = Pt(10)
    r_t = p_t.add_run("“Machine Learning-Driven Nano-QSAR and Quantum Chemical Design of Functionalized 2D Borophene Nanosheets for Targeted Disaggregation of Pathological Tau Fibrils in Alzheimer's Disease”")
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(74, 20, 140)
    
    doc.add_paragraph(
        "for consideration for publication as a Full Research Article in the Beilstein Journal of Nanotechnology."
    )
    
    doc.add_paragraph(
        "The study integrates GFN2-xTB quantum-chemical single-point interaction energies of 29 Alzheimer's / Tau "
        "therapeutics on a pristine beta-12 borophene cluster, physical AutoDock Vina docking against the human "
        "Tau filament cryo-EM structure, and a leak-free cross-validated surrogate model. A chi3-PEG-Tf "
        "functionalized allotrope is discussed only as future work (no real structural or quantum data for it "
        "exist). Three cationic phenothiazine dyes show anomalous positive interaction energies from steric clash "
        "and are reported as such and flagged out-of-domain; every value in the manuscript is computed from the "
        "deposited pipeline."
    )
    
    doc.add_paragraph(
        "All authors have approved the manuscript and confirm no competing interests."
    )
    
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(14)
    p_sign.add_run(
        "Sincerely,\n\n"
        "Andrés Monreal Hernández, Ph.D. (Corresponding Author)\n"
        "Universidad Estatal de Sonora, Mexico\n"
        "Email: andres.monreal@ues.mx"
    )
    
    out_docx = os.path.join(sub_dir, "01_Cover_Letter_Beilstein_Tau.docx")
    doc.save(out_docx)
    print(f"Generated Tau Cover Letter: {out_docx}")


def create_tau_cover_letter_md(sub_dir):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)
    fo = doc.styles['Normal'].font
    fo.name = 'Times New Roman'; fo.size = Pt(11); fo.color.rgb = RGBColor(33, 33, 33)
    doc.add_paragraph("Andrés Monreal Hernández, Ph.D.\nUniversidad Estatal de Sonora, Hermosillo, Sonora, Mexico\n"
                      "Email: andres.monreal@ues.mx | ORCID: 0009-0009-1207-8597").runs[0].font.bold = True
    doc.add_paragraph("To: The Editor-in-Chief, Molecular Diversity (Springer Nature)")
    doc.add_paragraph("Dear Editor,")
    doc.add_paragraph("We submit our original research manuscript for consideration in Molecular Diversity:")
    r = doc.add_paragraph().add_run("“Machine Learning-Driven Nano-QSAR and Quantum Chemical Design of Functionalized "
                                    "2D Borophene Nanosheets for Targeted Disaggregation of Pathological Tau Fibrils "
                                    "in Alzheimer's Disease”")
    r.font.bold = True; r.font.color.rgb = RGBColor(74, 20, 140)
    doc.add_paragraph("Real, pipeline-traceable results:").runs[0].font.bold = True
    for h in [
        "GFN2-xTB single-point interaction energies for 29 Alzheimer's / Tau therapeutics on a pristine beta-12 "
        "borophene cluster; frontier-orbital and conceptual-DFT indices taken from the xtb output.",
        "Physical AutoDock Vina v1.2.7 docking against the human Tau filament cryo-EM structure.",
        "Leak-free nested 5x5 cross-validated surrogate model on the real interaction energies; the "
        "feature-importance analysis is presented as exploratory.",
        "Three cationic phenothiazine dyes show anomalous positive interaction energies (steric clash) and are "
        "reported as such and flagged out-of-domain.",
        "The chi3-PEG-Tf functionalized allotrope is proposed as future work; it has no real data in this study.",
        "Full open-source pipeline and data archive (Zenodo 10.5281/zenodo.22187835).",
    ]:
        p = doc.add_paragraph(h); p.paragraph_format.left_indent = Inches(0.3)
    doc.add_paragraph("The manuscript is original, not under consideration elsewhere, and all authors approve the "
                      "submission and declare no competing interests.")
    doc.add_paragraph("Sincerely,\nAndrés Monreal Hernández, Ph.D. (Corresponding Author)")
    doc.save(os.path.join(sub_dir, "01_Cover_Letter_Molecular_Diversity.docx"))
    print("Generated Tau Molecular Diversity Cover Letter")


def build_tau_submission_bundle():
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    sub_dir = os.path.join(base_dir, "manuscript", "submission_ready")
    os.makedirs(sub_dir, exist_ok=True)
    
    create_tau_cover_letter(sub_dir)
    create_tau_cover_letter_md(sub_dir)
    
    src_docx = os.path.join(base_dir, "manuscript", "Beilstein_Manuscript_Tau_Borophene_Monreal_Hernandez_et_al.docx")
    dst_docx = os.path.join(sub_dir, "02_Main_Manuscript_Tau_Borophene_Monreal_Hernandez_et_al.docx")
    if os.path.exists(src_docx):
        shutil.copyfile(src_docx, dst_docx)

    src_si = os.path.join(base_dir, "manuscript", "Tau_Borophene_Supporting_Information.docx")
    dst_si = os.path.join(sub_dir, "03_Supporting_Information_Tau_Borophene_Monreal_Hernandez_et_al.docx")
    if os.path.exists(src_si):
        shutil.copyfile(src_si, dst_si)
        
    zip_path = os.path.join(base_dir, "borophene-alzheimer-tau-ai-FINAL-SUBMISSION-READY.zip")
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_f:
        for root, dirs, files in os.walk(sub_dir):
            for file in files:
                file_path = os.path.join(root, file)
                rel_path = os.path.relpath(file_path, sub_dir)
                zip_f.write(file_path, os.path.join("submission_ready", rel_path))
                
    print(f"\n=======================================================")
    print(f">>> TAU SUBMISSION PACKAGE GENERATED SUCCESSFULLY ({os.path.getsize(zip_path)} bytes) <<<")
    print(f" -> {zip_path}")
    print(f"=======================================================")

if __name__ == "__main__":
    build_tau_submission_bundle()
